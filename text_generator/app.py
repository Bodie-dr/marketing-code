import logging
import tempfile
import os
import sys
from pathlib import Path


if __package__ in {None, ""}:
    sys.path.insert(
        0,
        str(Path(__file__).resolve().parent.parent),
    )

import gradio as gr
from docx import Document
from docx.shared import Inches
from database import create_connection, create_schema
from config import DOCUMENTS_FOLDER
from document_repository import get_or_create_project, get_project_names

from text_generation import generate_text


def initialize_database() -> None:
    connection = create_connection()
    try:
        create_schema(connection)
    finally:
        connection.close()


logger = logging.getLogger(__name__)


def laad_bedrijven():
    connection = None

    try:
        connection = create_connection()
        create_schema(connection)
        bedrijven = get_project_names(connection)
        if not bedrijven and DOCUMENTS_FOLDER.is_dir():
            for bedrijfsmap in sorted(DOCUMENTS_FOLDER.iterdir()):
                if bedrijfsmap.is_dir():
                    get_or_create_project(connection, bedrijfsmap.name)
            connection.commit()
            bedrijven = get_project_names(connection)
    except Exception:
        logger.exception("Bedrijven konden niet uit de database worden geladen.")
        bedrijven = []
    finally:
        if connection is not None:
            connection.close()

    return gr.Dropdown(
        choices=bedrijven,
        value=bedrijven[0] if bedrijven else None,
    )


def save_generated_word(text):
    if not text or not text.strip():
        return None

    output_dir = Path(__file__).resolve().parent / "outputs"
    output_dir.mkdir(exist_ok=True)

    with tempfile.NamedTemporaryFile(
        suffix=".docx",
        dir=output_dir,
        delete=False,
    ) as bestand:
        bestand_path = Path(bestand.name)

    document = Document()
    image_path = Path(__file__).resolve().parent / "image.jpg"
    if not image_path.exists():
        image_path = Path(__file__).resolve().parent / "images.jpg"
    if image_path.exists():
        document.add_picture(str(image_path), width=Inches(6.5))

    for regel in text.strip().splitlines():
        document.add_paragraph(regel)
    document.save(bestand_path)
    return str(bestand_path)


def validate_and_generate(
    document,
    style_text,
    prompt,
    aanleiding,
    insteek,
    doelgroep,
    bedrijf,
    kanaal,
):
    opdracht = "\n".join(
        waarde
        for waarde in [aanleiding, insteek, doelgroep, bedrijf]
        if waarde and waarde.strip()
    )

    if not opdracht:
        raise gr.Error(
            "Vul minimaal de aanleiding, insteek, doelgroep of het bedrijf in."
        )

    try:
        resultaat = generate_text(
            opdracht.strip(),
            document=document,
            style_text=style_text or "",
            prompt=prompt or "",
            modus="Nieuwe tekst genereren",
            aanleiding=aanleiding or "",
            insteek=insteek or "",
            doelgroep=doelgroep or "",
            bedrijf=bedrijf or "",
            kanaal=kanaal or "LinkedIn",
        )
        return resultaat, save_generated_word(resultaat)
    except Exception as error:
        raise gr.Error(
            "De factuurtekst kon niet worden gegenereerd. "
            "Controleer je API-configuratie en probeer opnieuw."
        ) from error


APP_THEME = gr.themes.Base(
    primary_hue="cyan",
    secondary_hue="blue",
    neutral_hue="slate",
)

APP_CSS = """
:root {
    --tvb-blue: #222D4F;
    --tvb-blue-dark: #1A2340;
    --tvb-green: #38B5A8;
    --tvb-green-dark: #2D9A8E;
    --tvb-light: #F5F7FA;
    --tvb-secondary: #B8C4D2;
    --tvb-white: #FFFFFF;
}
/* Tekst boven en rondom invoervelden wit maken */
.gradio-container label,
.gradio-container label span,
.gradio-container .label-wrap,
.gradio-container .info,
.gradio-container p {
    color: #FFFFFF !important;
}

/* Markdown tekst op de donkerblauwe achtergrond */
.gradio-container .prose,
.gradio-container .prose p {
    color: #FFFFFF !important;
}

/* Tekst IN invoervelden juist donker houden */
.gradio-container input,
.gradio-container textarea,
.gradio-container select {
    color: #222D4F !important;
}

/* Placeholder in invoervelden */
.gradio-container input::placeholder,
.gradio-container textarea::placeholder {
    color: #6B7280 !important;
}

.gradio-container {
    background: linear-gradient(
        135deg,
        #222D4F 0%,
        #2A3763 35%,
        #F5F7FA 100%
    ) !important;
}

h1 {
    color: var(--tvb-blue) !important;
    font-size: 3.5rem !important;
    font-weight: 900 !important;
    text-align: center !important;
    border-bottom: 4px solid var(--tvb-green);
    padding-bottom: 10px;
    margin-bottom: 20px;
}

h2,h3,h4,h5,h6 {
    color: var(--tvb-blue) !important;
}

/* Kaarten */
.panel,
.gr-group,
.gr-box {
    background: white !important;
    border: 2px solid var(--tvb-secondary) !important;
    border-radius: 18px !important;
    box-shadow: 0 8px 25px rgba(34,45,79,0.15) !important;
}

/* Tabs */
button.selected {
    background: var(--tvb-blue) !important;
    color: white !important;
    border: 2px solid var(--tvb-green) !important;
}

button.selected:hover {
    background: var(--tvb-blue-dark) !important;
}


/* Primaire knoppen */
button.primary,
button[variant="primary"] {
    background: linear-gradient(
        90deg,
        var(--tvb-blue),
        #304064
    ) !important;
    color: white !important;
    border: none !important;
    border-radius: 10px !important;
    font-weight: 700 !important;
    box-shadow: 0 4px 12px rgba(34,45,79,.3) !important;
}


/* Hover */
button.primary:hover,
button[variant="primary"]:hover {
    background: var(--tvb-blue-dark) !important;
}

/* Secondary */
button.secondary,
button[variant="secondary"] {
    background: var(--tvb-green) !important;
    color: white !important;
    border: none !important;
    border-radius: 10px !important;
    font-weight: 700 !important;
}

button.secondary:hover,
button[variant="secondary"]:hover {
    background: var(--tvb-green-dark) !important;
}

/* Radio buttons */
input[type="radio"] {
    accent-color: var(--tvb-green) !important;
}

/* Upload component */
[data-testid="file-upload"] {
    border: 2px dashed var(--tvb-green) !important;
    border-radius: 12px !important;
}

/* Inputs */
input,
textarea,
select {
    border: 2px solid var(--tvb-secondary) !important;
    border-radius: 10px !important;
    background: white !important;
}

input:focus,
textarea:focus,
select:focus {
    border-color: var(--tvb-green) !important;
    box-shadow: 0 0 0 3px rgba(56,181,168,.2) !important;
}
/* Voorkom witte randen bij scrollen */
html,
body {
    margin: 0 !important;
    padding: 0 !important;
    background-color: #222D4F !important;
    min-height: 100% !important;
}

body {
    min-height: 100vh !important;
}

/* Laat Gradio de volledige pagina bedekken */
.gradio-container {
    width: 100% !important;
    max-width: none !important;
    min-height: 100vh !important;
    margin: 0 !important;
    background-color: #222D4F !important;
}
/* =========================================
   FIX: WITTE ZIJKANTEN BIJ SCROLLEN
   ========================================= */

html,
body,
gradio-app,
.gradio-container,
.main,
main {
    background: #222D4F !important;
    background-color: #222D4F !important;
}

/* Hele browserbreedte gebruiken */
html,
body {
    width: 100% !important;
    min-width: 100% !important;
    min-height: 100vh !important;
    margin: 0 !important;
    padding: 0 !important;
    overflow-x: hidden !important;
}

/* Gradio root volledig vullen */
gradio-app {
    display: block !important;
    width: 100% !important;
    min-height: 100vh !important;
    margin: 0 !important;
    padding: 0 !important;
}

/* Gradio container geen witte buitenruimte geven */
.gradio-container {
    width: 100vw !important;
    max-width: 100vw !important;
    min-height: 100vh !important;
    margin: 0 !important;
    box-sizing: border-box !important;
}
/* Bovenste titelkaart */
#factuur-title {
    background: #FFFFFF !important;
    border: 2px solid #B8C4D2 !important;
    border-left: 7px solid #38B5A8 !important;
    border-radius: 16px !important;
    padding: 22px 26px !important;
    margin-bottom: 18px !important;
    box-shadow: 0 8px 25px rgba(34, 45, 79, 0.18) !important;
}

/* Alle tekst in de bovenste titelkaart donker houden */
#factuur-title,
#factuur-title h1,
#factuur-title h2,
#factuur-title h3,
#factuur-title p,
#factuur-title .prose,
#factuur-title .prose p {
    color: #222D4F !important;
}

/* Hoofdtitel */
#factuur-title h1 {
    color: #222D4F !important;
    font-size: 3rem !important;
    font-weight: 900 !important;
    text-align: center !important;
    border-bottom: 4px solid #38B5A8 !important;
    padding-bottom: 12px !important;
    margin-top: 0 !important;
    margin-bottom: 14px !important;
}

/* Beschrijving onder de hoofdtitel */
#factuur-title p {
    color: #222D4F !important;
    text-align: center !important;
    font-size: 1.05rem !important;
    margin-bottom: 0 !important;
}

/* Titelkaart in de Factuurgenerator-tab */
#originele-tekst-title {
    background: #FFFFFF !important;
    border: 2px solid #B8C4D2 !important;
    border-left: 7px solid #38B5A8 !important;
    border-radius: 14px !important;
    padding: 16px 20px !important;
    margin-top: 14px !important;
    margin-bottom: 18px !important;
    box-shadow: 0 6px 18px rgba(34, 45, 79, 0.15) !important;
}

/* Tekst van de titel donkerblauw houden */
#originele-tekst-title,
#originele-tekst-title h1,
#originele-tekst-title h2,
#originele-tekst-title h3,
#originele-tekst-title p,
#originele-tekst-title .prose,
#originele-tekst-title .prose p {
    color: #222D4F !important;
}

/* Opmaak van de tweede titel */
#originele-tekst-title h2 {
    color: #222D4F !important;
    font-size: 1.55rem !important;
    font-weight: 800 !important;
    margin: 0 !important;
    border: none !important;
}
"""

with gr.Blocks(
    title="FotoModel & Factuurgenerator",
) as demo:

    gr.Markdown(
        """
    # Factuurgeneratie

    Genereer automatisch marketingteksten en download ze als Word-document.
    """,
        elem_id="factuur-title",
    )

    # ==================================================
    # TAB 1 - FACTUURGENERATOR
    # ==================================================

    with gr.Tab("Factuurgenerator"):

        gr.Markdown(
            """
        # Originele tekst maken in de stijl van de trainingsteksten
        """,
            elem_id="originele-tekst-title",
        )

        document = gr.File(
            label="Originele tekst of opdracht voor de AI (optioneel) ",
            file_types=[".txt", ".md", ".csv", ".pdf", ".docx"],
            type="filepath",
        )

        style_text = gr.Textbox(
            label="Originele tekst of opdracht voor de AI",
            lines=6,
            placeholder=(
                "Plak hier een voorbeeldtekst of stijlregels, bijv.: korte zinnen, "
                "formele toon, veel concrete termen."
            ),
        )

        tekst_prompt = gr.Textbox(
            label="Tekstbewerking",
            lines=4,
            placeholder=(
                "Bijvoorbeeld: schrijf kort en gebruik maximaal 5 opsommingstekens."
            ),
        )

        with gr.Row():
            aanleiding = gr.Textbox(
                label="Aanleiding",
                lines=3,
                placeholder="Waarom wordt deze tekst gemaakt?",
            )

            insteek = gr.Textbox(
                label="Insteek",
                lines=3,
                placeholder="Welke invalshoek of boodschap moet centraal staan?",
            )

        with gr.Row():
            doelgroep = gr.Textbox(
                label="Doelgroep",
                lines=3,
                placeholder="Voor wie is deze tekst bedoeld?",
            )

            bedrijf = gr.Dropdown(
                label="Bedrijf",
                choices=[],
                allow_custom_value=False,
                info="De AI zoekt de tone of voice en kernwaarden van dit bedrijf.",
            )

        kanaal = gr.Dropdown(
            choices=["LinkedIn", "Website", "Instagram"],
            value="LinkedIn",
            label="Kanaal",
        )

        genereer_button = gr.Button(
            "Tekst maken",
            variant="primary",
        )

        wis_button = gr.Button("Omschrijving wissen", variant="secondary")

        factuur_resultaat = gr.Textbox(
            label="Gegenereerde factuurtekst",
            lines=12,
        )

        download_file = gr.File(
            label="Download als Word-document",
            type="filepath",
        )

        genereer_button.click(
            fn=validate_and_generate,
            inputs=[
                document,
                style_text,
                tekst_prompt,
                aanleiding,
                insteek,
                doelgroep,
                bedrijf,
                kanaal,
            ],
            outputs=[factuur_resultaat, download_file],
        )

        wis_button.click(
            fn=lambda: (None, "", "", "", "", "", None, "LinkedIn", None),
            inputs=[],
            outputs=[
                document,
                style_text,
                tekst_prompt,
                aanleiding,
                insteek,
                doelgroep,
                bedrijf,
                kanaal,
                download_file,
            ],
        )

    demo.load(
        fn=laad_bedrijven,
        inputs=[],
        outputs=bedrijf,
    )

    # ==================================================
    # LEGACY BUTTON HANDLERS
    # ==================================================

    # Kept intentionally empty to avoid duplicate callbacks after the tab reorder.



if __name__ == "__main__":
    initialize_database()

    port = int(os.environ.get("PORT", 7860))

    demo.launch(
        server_name="0.0.0.0",
        server_port=port,
        theme=APP_THEME,
        css=APP_CSS,
    )

