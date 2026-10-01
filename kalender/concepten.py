"""Maak per geplande post een Word-concept met de tekstgenerator."""

import logging
import re
import sys
from dataclasses import dataclass
from pathlib import Path

from docx import Document

from planning_posts import PROJECT_DIR, Post, formatteer_datum


logger = logging.getLogger(__name__)

CONCEPTEN_DIR = PROJECT_DIR / "outputs" / "concepten"

# Planningskanaal -> kanaal waarop de tekstgenerator voorbeeldteksten zoekt.
GENERATOR_KANAAL = {
    "Socials": "LinkedIn",
    "LinkedIn": "LinkedIn",
    "Instagram": "Instagram",
    "LinkedIn + Instagram": "LinkedIn",
    "TikTok": "Instagram",
    "Nieuwsbrief": "Website",
    "Website": "Website",
}

KANAAL_INSTRUCTIE = {
    "Socials": "Schrijf een social-media post die op LinkedIn en Instagram geplaatst kan worden.",
    "LinkedIn": "Schrijf een LinkedIn-post.",
    "Instagram": "Schrijf een Instagram-caption: kort, persoonlijk, met passende hashtags.",
    "LinkedIn + Instagram": "Schrijf een post die zowel op LinkedIn als Instagram geplaatst kan worden.",
    "TikTok": (
        "Schrijf een TikTok-concept: een pakkende hook voor de eerste seconden, "
        "een korte opzet van de video in 3 tot 5 shots en een caption met hashtags."
    ),
    "Nieuwsbrief": "Schrijf een nieuwsbriefartikel met een korte kop en inleiding.",
    "Website": "Schrijf een tekst voor de website.",
}

PLACEHOLDER_INSTRUCTIE = (
    "De planning noemt alleen het thema. Vul ontbrekende details niet zelf in, "
    "maar gebruik duidelijke placeholders tussen blokhaken, zoals [naam collega], "
    "[functie], [projectnaam] of [locatie], zodat de marketeer ze kan aanvullen."
)


@dataclass
class ConceptResultaat:
    post: Post
    pad: Path
    status: str  # "gemaakt", "bestond al" of "mislukt"
    fout: str = ""


# Achteraan toevoegen: modules uit deze map gaan voor, tekstgenerator vult aan.
TEXT_GENERATOR_DIR = str(PROJECT_DIR / "tekstgenerator")
if TEXT_GENERATOR_DIR not in sys.path:
    sys.path.append(TEXT_GENERATOR_DIR)


def _generate_text(**kwargs) -> str:
    """Laad de tekstgenerator pas als er echt een concept gemaakt wordt."""
    from text_generation import generate_text

    return generate_text(**kwargs)


def _veilige_naam(tekst: str, max_lengte: int = 60) -> str:
    tekst = re.sub(r'[<>:"/\\|?*\n\r\t]+', " ", tekst)
    tekst = re.sub(r"\s+", " ", tekst).strip(" .")
    return tekst[:max_lengte].strip() or "post"


def concept_pad(post: Post, output_dir: Path = CONCEPTEN_DIR) -> Path:
    jaar, week, _ = post.datum.isocalendar()
    map_naam = f"{jaar}-week-{week:02d}"
    onderwerp = post.onderwerp or post.soort_content or "post"
    bestandsnaam = _veilige_naam(f"{post.datum:%Y-%m-%d} {post.kanaal} - {onderwerp}") + ".docx"
    return output_dir / map_naam / _veilige_naam(post.bedrijf) / bestandsnaam


def generator_argumenten(post: Post) -> dict:
    insteek = " – ".join(deel for deel in [post.onderwerp, post.soort_content] if deel)
    aanleiding = f"Geplande post volgens de social-media jaarplanning ({post.datum_tekst})."
    if post.bijzonderheden:
        aanleiding += f" Bijzonderheden deze week: {post.bijzonderheden}."

    return {
        "nieuwe_opdracht": f"{post.bedrijf}\n{insteek}\n{aanleiding}",
        "aanleiding": aanleiding,
        "insteek": insteek,
        "doelgroep": "",
        "bedrijf": post.bedrijf,
        "kanaal": GENERATOR_KANAAL.get(post.kanaal, "LinkedIn"),
        "prompt": f"{KANAAL_INSTRUCTIE.get(post.kanaal, '')} {PLACEHOLDER_INSTRUCTIE}".strip(),
        "modus": "Nieuwe tekst genereren",
    }


def schrijf_word(post: Post, tekst: str, pad: Path) -> Path:
    document = Document()
    document.add_heading(f"{post.bedrijf} – {post.kanaal}", level=1)

    for label, waarde in [
        ("Datum", formatteer_datum(post.datum) if post.dag_bekend else post.datum_tekst),
        ("Onderwerp", post.onderwerp),
        ("Soort content", post.soort_content),
        ("Bijzonderheden", post.bijzonderheden),
    ]:
        if waarde:
            alinea = document.add_paragraph()
            alinea.add_run(f"{label}: ").bold = True
            alinea.add_run(waarde)

    status = document.add_paragraph()
    status.add_run("CONCEPT – nalopen en [placeholders] invullen vóór publicatie.").italic = True

    document.add_heading("Tekst", level=2)
    for regel in tekst.strip().splitlines():
        document.add_paragraph(regel)

    pad.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(pad))
    return pad


def maak_concepten(
    posts: list[Post],
    opnieuw: bool = False,
    output_dir: Path = CONCEPTEN_DIR,
    voortgang=None,
) -> list[ConceptResultaat]:
    """Maak concepten; bestaande concepten worden overgeslagen (dat scheelt API-kosten)."""
    resultaten = []

    for nummer, post in enumerate(posts, start=1):
        pad = concept_pad(post, output_dir)
        if voortgang:
            voortgang(nummer, len(posts), post)

        if pad.exists() and not opnieuw:
            logger.info("Bestaat al, overgeslagen: %s", pad.name)
            resultaten.append(ConceptResultaat(post, pad, "bestond al"))
            continue

        try:
            tekst = _generate_text(**generator_argumenten(post))
            schrijf_word(post, tekst, pad)
            logger.info("[%s/%s] Concept gemaakt: %s", nummer, len(posts), pad.name)
            resultaten.append(ConceptResultaat(post, pad, "gemaakt"))
        except Exception as error:
            logger.exception("Concept mislukt voor %s (%s)", post.bedrijf, post.datum_tekst)
            resultaten.append(ConceptResultaat(post, pad, "mislukt", str(error)))

    return resultaten
