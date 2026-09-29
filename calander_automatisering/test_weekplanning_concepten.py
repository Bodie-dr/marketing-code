"""Tests voor weekoverzicht, agenda-export en concepten (zonder API-aanroepen).

Start: cd calander_automatisering; python -m unittest test_weekplanning_concepten -v
"""

import tempfile
import unittest
from datetime import date, datetime, timezone
from pathlib import Path
from unittest.mock import patch

from docx import Document

import concepten
from agenda_export import posts_naar_ics
from planning_posts import Post
from weekplanning import (
    ALLE_BEDRIJVEN,
    bedrijven_in_planning,
    filter_bedrijf,
    maandag_van,
    posts_in_periode,
    posts_tabel,
    standaard_week,
    week_label,
    weken_in_planning,
)


def maak_post(datum, kanaal="Socials", dag_bekend=True, bron="TTI", **velden):
    waarden = {
        "bedrijf": "Terberg Totaal Installaties",
        "onderwerp": "Employer branding",
        "soort_content": "Carriere verhaal",
        "bijzonderheden": None,
    }
    waarden.update(velden)
    return Post(
        bron=bron,
        week=datum.isocalendar().week,
        datum=datum,
        dag_bekend=dag_bekend,
        kanaal=kanaal,
        **waarden,
    )


class WeekplanningTests(unittest.TestCase):
    def setUp(self):
        self.posts = [
            maak_post(date(2026, 10, 6)),   # di week 41
            maak_post(date(2026, 10, 8)),   # do week 41
            maak_post(date(2026, 10, 13)),  # di week 42
            maak_post(date(2026, 11, 3)),   # di week 45
        ]

    def test_periode_begint_op_maandag_van_de_week(self):
        # Ook vanaf een woensdag hoort de dinsdagpost van die week erbij.
        geselecteerd = posts_in_periode(self.posts, date(2026, 10, 7), 1)
        self.assertEqual([p.datum.day for p in geselecteerd], [6, 8])
        self.assertEqual(len(posts_in_periode(self.posts, date(2026, 10, 7), 2)), 3)

    def test_week_label_en_keuzes(self):
        self.assertEqual(maandag_van(date(2026, 10, 8)), date(2026, 10, 5))
        self.assertEqual(week_label(date(2026, 10, 5)), "Week 41 · 5 okt – 11 okt 2026")
        self.assertEqual([waarde for _, waarde in weken_in_planning(self.posts)],
                         ["2026-10-05", "2026-10-12", "2026-11-02"])

    def test_standaard_week_is_deze_of_volgende_week_met_posts(self):
        self.assertEqual(standaard_week(self.posts, date(2026, 10, 14)), "2026-10-12")
        self.assertEqual(standaard_week(self.posts, date(2026, 10, 20)), "2026-11-02")
        self.assertEqual(standaard_week(self.posts, date(2027, 1, 1)), "2026-11-02")
        self.assertIsNone(standaard_week([], date(2026, 1, 1)))

    def test_filter_op_bedrijf_of_werkblad(self):
        posts = self.posts + [maak_post(date(2026, 10, 7), bedrijf="TVB", bron="TVB")]

        self.assertEqual(len(filter_bedrijf(posts, "Terberg Totaal Installaties")), 4)
        self.assertEqual(len(filter_bedrijf(posts, "tti")), 4)
        self.assertEqual([p.bedrijf for p in filter_bedrijf(posts, " TVB ")], ["TVB"])
        self.assertEqual(filter_bedrijf(posts, "Onbekend"), [])
        self.assertEqual(filter_bedrijf(posts, None), posts)
        self.assertEqual(filter_bedrijf(posts, ALLE_BEDRIJVEN), posts)
        self.assertEqual(bedrijven_in_planning(posts), ["Terberg Totaal Installaties", "TVB"])

    def test_tabel(self):
        tabel = posts_tabel(self.posts[:1])
        self.assertEqual(tabel.loc[0, "Datum"], "di 6 okt 2026")
        self.assertEqual(list(posts_tabel([]).columns)[0], "Datum")


class AgendaTests(unittest.TestCase):
    def ics(self, posts):
        return posts_naar_ics(posts, nu=datetime(2026, 9, 29, tzinfo=timezone.utc))

    def test_hele_dag_afspraak_met_omschrijving(self):
        ics = self.ics([maak_post(date(2026, 10, 6), bijzonderheden="Dag van de monteur, inhaker")])

        self.assertTrue(ics.startswith("BEGIN:VCALENDAR\r\n"))
        self.assertTrue(ics.endswith("END:VCALENDAR\r\n"))
        self.assertIn("DTSTART;VALUE=DATE:20261006\r\n", ics)
        self.assertIn("DTEND;VALUE=DATE:20261007\r\n", ics)
        self.assertIn("SUMMARY:Terberg Totaal Installaties – Socials: Employer branding", ics)
        # Komma's escapen; lange regels worden gevouwen.
        self.assertIn("Dag van de monteur\\, inhaker", ics.replace("\r\n ", ""))

    def test_post_zonder_vaste_dag_beslaat_werkweek(self):
        ics = self.ics([maak_post(date(2026, 10, 5), kanaal="TikTok", dag_bekend=False)])

        self.assertIn("DTSTART;VALUE=DATE:20261005\r\n", ics)
        self.assertIn("DTEND;VALUE=DATE:20261010\r\n", ics)
        self.assertIn("(dag vrij)", ics.replace("\r\n ", ""))

    def test_regels_maximaal_75_bytes_en_unieke_uids(self):
        posts = [
            maak_post(date(2026, 10, 6), onderwerp="Project " * 30),
            maak_post(date(2026, 10, 6), onderwerp="Tweede post zelfde dag"),
        ]
        ics = self.ics(posts)

        self.assertTrue(all(len(regel.encode("utf-8")) <= 75 for regel in ics.split("\r\n")))
        uids = [regel for regel in ics.split("\r\n") if regel.startswith("UID:")]
        self.assertEqual(len(set(uids)), 2)
        self.assertEqual(ics, self.ics(posts), "UIDs moeten stabiel zijn tussen exports")


class ExcelExportTests(unittest.TestCase):
    def test_een_regel_per_post_met_echte_datum(self):
        from openpyxl import load_workbook

        from excel_export import schrijf_excel

        posts = [
            maak_post(date(2026, 10, 6), bijzonderheden="Inhaker"),
            maak_post(date(2026, 10, 5), kanaal="TikTok", dag_bekend=False, soort_content=None),
        ]
        with tempfile.TemporaryDirectory() as folder:
            pad = schrijf_excel(posts, Path(folder) / "sub" / "planning.xlsx")
            blad = load_workbook(pad).active
            rijen = list(blad.iter_rows(values_only=True))
            filter_bereik = blad.auto_filter.ref
            vastgezet = blad.freeze_panes

        self.assertEqual(rijen[0][:5], ("Datum", "Dag", "Week", "Bedrijf", "Kanaal"))
        self.assertEqual(rijen[1][0].date(), date(2026, 10, 6))
        self.assertEqual(rijen[1][1:5], ("di", 41, "Terberg Totaal Installaties", "Socials"))
        self.assertEqual(rijen[1][7], "Inhaker")
        self.assertEqual(rijen[2][1], "vrij")
        self.assertIsNone(rijen[2][6])  # lege soort content blijft leeg
        self.assertEqual(filter_bereik, "A1:I3")
        self.assertEqual(vastgezet, "A2")


class ConceptenTests(unittest.TestCase):
    def test_argumenten_voor_generator(self):
        argumenten = concepten.generator_argumenten(
            maak_post(date(2026, 10, 5), kanaal="TikTok", dag_bekend=False, bijzonderheden="Herfstvakantie")
        )

        self.assertEqual(argumenten["bedrijf"], "Terberg Totaal Installaties")
        self.assertEqual(argumenten["kanaal"], "Instagram")
        self.assertEqual(argumenten["insteek"], "Employer branding – Carriere verhaal")
        self.assertIn("Herfstvakantie", argumenten["aanleiding"])
        self.assertIn("TikTok", argumenten["prompt"])
        self.assertIn("[naam collega]", argumenten["prompt"])

    def test_elk_planningskanaal_heeft_generatorkanaal(self):
        for kanaal in ["Socials", "LinkedIn", "Instagram", "LinkedIn + Instagram", "TikTok", "Nieuwsbrief", "Website"]:
            with self.subTest(kanaal=kanaal):
                self.assertIn(concepten.GENERATOR_KANAAL[kanaal], {"LinkedIn", "Instagram", "Website"})
                self.assertIn(kanaal, concepten.KANAAL_INSTRUCTIE)

    def test_concept_pad_per_week_en_bedrijf(self):
        pad = concepten.concept_pad(
            maak_post(date(2026, 10, 6), onderwerp="Projecten / zakelijk"), Path("uit")
        )
        self.assertEqual(pad.parts[:3], ("uit", "2026-week-41", "Terberg Totaal Installaties"))
        self.assertEqual(pad.name, "2026-10-06 Socials - Projecten zakelijk.docx")

    def test_maakt_word_en_slaat_bestaande_over(self):
        posts = [maak_post(date(2026, 10, 6)), maak_post(date(2026, 10, 8))]
        with tempfile.TemporaryDirectory() as folder, patch.object(
            concepten, "_generate_text", return_value="Regel 1\nRegel 2"
        ) as generate:
            eerste = concepten.maak_concepten(posts, output_dir=Path(folder))
            tweede = concepten.maak_concepten(posts, output_dir=Path(folder))
            opnieuw = concepten.maak_concepten(posts[:1], output_dir=Path(folder), opnieuw=True)

            tekst = "\n".join(p.text for p in Document(eerste[0].pad).paragraphs)

        self.assertEqual([r.status for r in eerste], ["gemaakt", "gemaakt"])
        self.assertEqual([r.status for r in tweede], ["bestond al", "bestond al"])
        self.assertEqual([r.status for r in opnieuw], ["gemaakt"])
        self.assertEqual(generate.call_count, 3)
        self.assertIn("Soort content: Carriere verhaal", tekst)
        self.assertIn("CONCEPT", tekst)
        self.assertIn("Regel 2", tekst)

    def test_fout_bij_een_post_stopt_de_rest_niet(self):
        posts = [maak_post(date(2026, 10, 6)), maak_post(date(2026, 10, 8))]
        with tempfile.TemporaryDirectory() as folder, patch.object(
            concepten, "_generate_text", side_effect=[RuntimeError("API weg"), "Tekst"]
        ), patch.object(concepten.logger, "exception"):
            resultaten = concepten.maak_concepten(posts, output_dir=Path(folder))

        self.assertEqual([r.status for r in resultaten], ["mislukt", "gemaakt"])
        self.assertEqual(resultaten[0].fout, "API weg")


if __name__ == "__main__":
    unittest.main()
