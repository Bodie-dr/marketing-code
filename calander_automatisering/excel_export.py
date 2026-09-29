"""Exporteer posts als opgemaakt Excel-bestand (één regel per post)."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from planning_posts import KORTE_DAGEN, Post


KOLOMMEN = [
    ("Datum", 13),
    ("Dag", 7),
    ("Week", 7),
    ("Bedrijf", 32),
    ("Kanaal", 20),
    ("Onderwerp", 32),
    ("Soort content", 36),
    ("Bijzonderheden", 36),
    ("Status", 14),
]

KOP_VULLING = PatternFill("solid", fgColor="222D4F")
KOP_LETTER = Font(bold=True, color="FFFFFF")


def schrijf_excel(posts: list[Post], pad: Path) -> Path:
    werkboek = Workbook()
    blad = werkboek.active
    blad.title = "Planning"

    blad.append([naam for naam, _ in KOLOMMEN])
    for post in posts:
        blad.append([
            post.datum,
            KORTE_DAGEN[post.datum.isoweekday() - 1] if post.dag_bekend else "vrij",
            post.week,
            post.bedrijf,
            post.kanaal,
            post.onderwerp or "",
            post.soort_content or "",
            post.bijzonderheden or "",
            "",  # in te vullen door het team, bijv. concept / gepland / geplaatst
        ])

    for kolom, (_, breedte) in enumerate(KOLOMMEN, start=1):
        blad.column_dimensions[get_column_letter(kolom)].width = breedte
        kop = blad.cell(row=1, column=kolom)
        kop.fill = KOP_VULLING
        kop.font = KOP_LETTER
        kop.alignment = Alignment(vertical="center")

    for rij in blad.iter_rows(min_row=2):
        rij[0].number_format = "DD-MM-YYYY"
        for cel in rij:
            cel.alignment = Alignment(vertical="top", wrap_text=True)

    blad.freeze_panes = "A2"
    blad.auto_filter.ref = blad.dimensions

    pad = Path(pad)
    pad.parent.mkdir(parents=True, exist_ok=True)
    werkboek.save(str(pad))
    return pad
